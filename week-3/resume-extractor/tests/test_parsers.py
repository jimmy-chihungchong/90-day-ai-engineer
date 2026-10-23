import io

import pytest
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from resume_extractor.parsers import ParseError, parse_resume

BODY = [
    "Jane Doe", "jane@example.com | +1 555 0100",
    "EXPERIENCE", "Senior Engineer, Acme Corp (2019 - Present)",
    "SKILLS", "Python, SQL, Docker",
]


def make_pdf(lines, pages=1) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for _ in range(pages):
        y = 750
        for ln in lines:
            c.drawString(72, y, ln)
            y -= 20
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(lines, table=None, header=None) -> bytes:
    d = Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    for ln in lines:
        d.add_paragraph(ln)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_text_and_pages():
    r = parse_resume("a.pdf", make_pdf(BODY, pages=2))
    assert r.filetype == "pdf" and r.pages == 2
    assert "Jane Doe" in r.text and "Docker" in r.text


def test_docx_paragraphs_table_header_order():
    data = make_docx(BODY, table=[["Skills", "Python, Go"], ["Languages", "English"]],
                     header="HEADER-NAME")
    r = parse_resume("a.DOCX", data)  # extension is case-insensitive
    assert r.text.startswith("HEADER-NAME")
    assert "Skills | Python, Go" in r.text
    assert r.text.index("Docker") < r.text.index("Skills | Python, Go")


def test_docx_merged_cells_not_duplicated():
    d = Document()
    d.add_paragraph("x" * 60)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged Heading"
    buf = io.BytesIO(); d.save(buf)
    assert parse_resume("a.docx", buf.getvalue()).text.count("Merged Heading") == 1


def test_image_only_pdf_rejected():
    with pytest.raises(ParseError, match="scanned"):
        parse_resume("scan.pdf", make_pdf([]))


@pytest.mark.parametrize("name", ["cv.doc", "cv.txt", "cv"])
def test_unsupported_types(name):
    with pytest.raises(ParseError):
        parse_resume(name, b"whatever")


def test_corrupt_files():
    with pytest.raises(ParseError, match="PDF"):
        parse_resume("a.pdf", b"not a pdf")
    with pytest.raises(ParseError, match="DOCX"):
        parse_resume("a.docx", b"not a docx")
